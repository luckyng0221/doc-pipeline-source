from __future__ import annotations

import json
import re
import threading
import time
from pathlib import Path

import pymupdf
import pytest
import yaml
from openpyxl import load_workbook

from doc_pipeline.config import load_config
from doc_pipeline.errors import ProviderError, RunStopped
from doc_pipeline.models import ModelConfig
from doc_pipeline.processor import BatchProcessor
from doc_pipeline.providers import ProviderResponse


USAGE = {"prompt_tokens": 10, "output_tokens": 2, "thoughts_tokens": 1, "cached_tokens": 0, "total_tokens": 12}
# Page text the boundary rules cannot decide, so page 2 goes to the boundary model.
UNMARKED_PAGES = ["Invoice 1 page 1", "Invoice 1 page 2"]
# Page text that settles the boundary and prints the extracted values.
MARKED_PAGES = ["Acme Trading\nPage 1 of 2\nTotal due 12.50", "Acme Trading delivery details\nPage 2 of 2"]


class ScriptedProvider:
    """Offline provider: answers boundary prompts, then extraction prompts.

    Extraction responses come from a per-model-name list of field payloads.
    """

    def __init__(self, model_name: str, payloads: list[dict]) -> None:
        self.model_name = model_name
        self.payloads = list(payloads)
        self.calls = 0
        self.config = ModelConfig(
            name=model_name, provider="gemini", model="test-model", api_key_env="X_API_KEY"
        )

    def generate_json(self, prompt: str, parts: list[bytes]) -> ProviderResponse:
        self.calls += 1
        if "boundaries between" in prompt:
            # Page 1 always starts a document; only page 2 is a candidate.
            return ProviderResponse(
                data={"pages": [{"page": 2, "new_document": False, "confidence": 0.99}]},
                usage=USAGE,
            )
        payload = self.payloads.pop(0)
        return ProviderResponse(data={"fields": payload}, usage=USAGE)


def _write_config(root: Path, **settings: object) -> Path:
    config = {
        "root": str(root),
        "paths": {
            "processing": "processing",
            "split": "split",
            "done": "done",
            "errors": "errors",
            "output": "output",
            "registry": "registry.db",
        },
        "boundary_model": {
            "name": "boundary",
            "provider": "gemini",
            "model": "test-model",
            "api_key_env": "X_API_KEY",
            "structured_output": True,
        },
        "models": [
            {
                "name": "primary_a",
                "provider": "gemini",
                "model": "test-model",
                "api_key_env": "X_API_KEY",
                "role": "primary",
            },
            {
                "name": "primary_b",
                "provider": "gemini",
                "model": "test-model",
                "api_key_env": "X_API_KEY",
                "role": "primary",
            },
            {
                "name": "arbiter",
                "provider": "gemini",
                "model": "test-model",
                "api_key_env": "X_API_KEY",
                "role": "arbiter",
            },
        ],
        "fields": [
            {"name": "vendor", "type": "string", "required": True},
            {"name": "total", "type": "decimal", "required": True},
        ],
        "fingerprint_fields": ["vendor", "total"],
        **settings,
    }
    path = root / "config.yaml"
    path.write_text(yaml.safe_dump(config), encoding="utf-8")
    return path


def _write_pdf(
    root: Path,
    texts: list[str] = UNMARKED_PAGES,
    hidden: list[str] | None = None,
    ocr: bool = False,
    low_opacity: bool = False,
) -> Path:
    """Write the PDF, planting one invisible line per page.

    hidden carries one string per page (empty plants nothing), drawn in
    invisible render mode 3 - or at a fill opacity of 0.05 when low_opacity
    is set. With ocr, the page first gets an image covering all of it, so the
    invisible line is an OCR-style layer over a scan and is treated as the
    page's own text.
    """
    source = root / "uploads" / "sample.pdf"
    source.parent.mkdir(parents=True, exist_ok=True)
    document = pymupdf.open()
    for text, ghost in zip(texts, hidden or [""] * len(texts)):
        page = document.new_page()
        if ocr:
            pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 595, 842), False)
            pixmap.clear_with(255)
            page.insert_image(page.rect, pixmap=pixmap)
        if text:
            page.insert_text((72, 72), text)
        if ghost and low_opacity:
            page.insert_text((72, 500), ghost, fill_opacity=0.05)
        elif ghost:
            page.insert_text((72, 500), ghost, render_mode=3)
    document.save(source)
    document.close()
    return source


def _processor_with_providers(
    root: Path,
    monkeypatch,
    settings: dict | None = None,
    stop=None,
    **payloads: list[dict],
) -> tuple[BatchProcessor, dict[str, ScriptedProvider]]:
    config = load_config(_write_config(root, **(settings or {})))
    processor = BatchProcessor(config, stop_check=stop)
    providers = {name: ScriptedProvider(name, list(payload)) for name, payload in payloads.items()}
    providers.setdefault("boundary", ScriptedProvider("boundary", []))

    def fake_provider(self, model: ModelConfig):
        return providers[model.name]

    monkeypatch.setattr(BatchProcessor, "_provider", fake_provider)
    return processor, providers


def _document_row(output: Path) -> dict:
    workbook = load_workbook(output, data_only=True)
    header, row = list(workbook["Documents"].iter_rows(values_only=True))[:2]
    return dict(zip(header, row))


def test_agreeing_primaries_skip_the_arbiter(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme", "total": "12.50"}],
        primary_b=[{"vendor": "Acme", "total": "12.50"}],
        arbiter=[{"vendor": "Acme", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path)

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None and output.is_file()
    workbook = load_workbook(output, data_only=True)
    row = _document_row(output)
    # The values are not printed in the PDF text, so the second primary runs.
    assert row["Validation status"] == "CONSENSUS_2_PRIMARIES"
    assert row["Models called"] == 2
    assert "Second model called because primary_a's answer could not be confirmed" in row["Remarks"]
    audit_rows = list(workbook["Model audit"].iter_rows(values_only=True))
    models = [str(row[audit_rows[0].index("Model")] or "") for row in audit_rows[1:]]
    assert not any("arbiter" in model for model in models)
    assert any("primary_a" in model for model in models)
    assert any("primary_b" in model for model in models)
    usage_rows = list(workbook["Token usage"].iter_rows(values_only=True))
    # Boundary (1 call) + two primaries (1 call each) = three usage rows.
    assert len(usage_rows) == 4  # header + 3 calls
    run_rows = dict(
        (row[0], row[1]) for row in workbook["Run info"].iter_rows(values_only=True)
    )
    assert any(str(key).startswith("Tokens ") for key in run_rows)
    assert run_rows["Pages sent to the boundary model"] == "1"


def test_disagreeing_primaries_call_the_arbiter(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme", "total": "12.50"}],
        primary_b=[{"vendor": "Other", "total": "12.50"}],
        arbiter=[{"vendor": "Acme", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path)

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None and output.is_file()
    row = _document_row(output)
    # Two of three models agree on every populated field.
    assert row["Validation status"] == "CONSENSUS_2_OF_3"
    assert row["vendor"] == "Acme"
    assert "Arbiter called because the primaries disagreed" in row["Remarks"]


def test_first_primary_confirmed_by_pdf_text_needs_no_other_model(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
        primary_b=[{"vendor": "Acme Trading", "total": "12.50"}],
        arbiter=[{"vendor": "Acme Trading", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path, MARKED_PAGES)

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "TEXT_VERIFIED"
    assert row["Models called"] == 1
    assert (row["Page start"], row["Page end"]) == (1, 2)
    assert row["Remarks"].startswith("Starts at page 1: first page")
    assert [provider.calls for provider in providers.values()] == [1, 0, 0, 0]
    workbook = load_workbook(output, data_only=True)
    pages = list(workbook["Pages"].iter_rows(values_only=True))
    assert [page[1:4] for page in pages[1:]] == [(1, "start", "rule"), (2, "continuation", "rule")]
    run_rows = dict((row[0], row[1]) for row in workbook["Run info"].iter_rows(values_only=True))
    assert run_rows["Boundary model requests"] == "0"
    assert run_rows["Documents confirmed by the PDF text (one model)"] == "1"


def test_unconfirmed_first_primary_brings_in_the_second(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme Trading", "total": "99.99"}],
        primary_b=[{"vendor": "Acme Trading", "total": "12.50"}],
        arbiter=[{"vendor": "Acme Trading", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path, MARKED_PAGES)

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "CONSENSUS_2_OF_3"
    assert row["total"] == 12.5
    assert "total was not found in the PDF text" in row["Remarks"]
    assert providers["arbiter"].calls == 1


def test_always_mode_runs_both_primaries(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        settings={"second_opinion": "always"},
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
        primary_b=[{"vendor": "Acme Trading", "total": "12.50"}],
        arbiter=[],
    )
    source = _write_pdf(tmp_path, MARKED_PAGES)

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    assert _document_row(output)["Validation status"] == "CONSENSUS_2_PRIMARIES"
    assert providers["primary_b"].calls == 1


ONE_MODEL = {
    "models": [
        {"name": "primary_a", "provider": "gemini", "model": "test-model", "api_key_env": "X_API_KEY"},
    ],
}

TWO_PRIMARIES = {
    "models": [
        {"name": "primary_a", "provider": "gemini", "model": "test-model", "api_key_env": "X_API_KEY"},
        {"name": "primary_b", "provider": "gemini", "model": "test-model", "api_key_env": "X_API_KEY"},
    ],
}


def test_rules_only_needs_no_model_and_says_so(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        settings={"models": [], "boundary_model": None},
    )
    source = _write_pdf(tmp_path, ["Acme Trading\nPage 1 of 1\nVendor: Acme Trading\nTotal: 12.50"])

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    row = _document_row(outcome.output_path)
    assert row["Validation status"] == "RULES_ONLY"
    assert row["Models called"] == 0
    assert row["vendor"] == "Acme Trading"
    assert row["total"] == 12.5
    assert row["Remarks"].startswith("Starts at page 1: first page | No AI model")
    assert providers["boundary"].calls == 0
    run_rows = dict(load_workbook(outcome.output_path, data_only=True)["Run info"].iter_rows(values_only=True))
    assert run_rows["Mode"] == "Rules only (no AI model); page splitting: rules only"


def test_single_model_is_verified_by_text_or_flagged(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        settings=ONE_MODEL,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}, {"vendor": "Acme Trading", "total": "99.99"}],
    )
    first = _write_pdf(tmp_path, MARKED_PAGES)
    verified = processor.handle_claimed_pdf(first)
    second = _write_pdf(tmp_path, [*MARKED_PAGES[:1], "Different delivery details\nPage 2 of 2"])
    flagged = processor.handle_claimed_pdf(second)

    assert _document_row(verified.output_path)["Validation status"] == "TEXT_VERIFIED"
    row = _document_row(flagged.output_path)
    assert row["Validation status"] == "SINGLE_MODEL"
    assert "total was not found in the PDF text" in row["Remarks"]


def test_hidden_text_keeps_confirmed_values_from_text_verified(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
        primary_b=[{"vendor": "Acme Trading", "total": "12.50"}],
        arbiter=[{"vendor": "Acme Trading", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path, MARKED_PAGES, hidden=["Invisible 99.99", ""])

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    # The values match the visible text, but hidden text could hold the rest,
    # so the second primary is asked and the document is not TEXT_VERIFIED.
    assert row["Validation status"] == "CONSENSUS_2_PRIMARIES"
    assert row["Models called"] == 2
    assert "Hidden text found on page 1; check values against the printed document" in row["Remarks"]
    assert "the PDF has hidden text on page 1" in row["Remarks"]
    assert providers["arbiter"].calls == 0


def test_an_ocr_layer_still_verifies_the_values_with_one_model(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(tmp_path, monkeypatch, settings=ONE_MODEL, primary_a=[{"vendor": "Acme Trading", "total": "12.50"}])
    # The total is printed only in the invisible layer, but a scan's image
    # covers the page, so the layer is the page's own text and can verify.
    source = _write_pdf(
        tmp_path,
        ["Acme Trading\nPage 1 of 2", "Acme Trading delivery details\nPage 2 of 2"],
        hidden=["Total due 12.50", ""],
        ocr=True,
    )

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "TEXT_VERIFIED"
    assert row["Models called"] == 1
    assert "Hidden text found" not in row["Remarks"]


def test_rules_only_reads_the_visible_total_and_reports_the_hidden_text(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        settings={"models": [], "boundary_model": None},
    )
    source = _write_pdf(
        tmp_path,
        ["Acme Trading\nPage 1 of 1\nVendor: Acme Trading\nTotal: 12.50"],
        hidden=["Total: 99.99"],
    )

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    row = _document_row(outcome.output_path)
    assert row["Validation status"] == "RULES_ONLY"
    assert row["total"] == 12.5
    assert "Hidden text found on page 1; check values against the printed document" in row["Remarks"]


def test_a_value_only_in_a_low_opacity_span_is_never_text_verified(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
        primary_b=[{"vendor": "Acme Trading", "total": "12.50"}],
        arbiter=[{"vendor": "Acme Trading", "total": "12.50"}],
    )
    # The total is printed only at a fill opacity of 0.05, under the visible
    # threshold: hidden text confirms nothing.
    source = _write_pdf(
        tmp_path,
        ["Acme Trading\nPage 1 of 2", "Acme Trading delivery details\nPage 2 of 2"],
        hidden=["Total due 12.50", ""],
        low_opacity=True,
    )

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "CONSENSUS_2_PRIMARIES"
    assert row["Models called"] == 2
    assert "total was not found in the PDF text" in row["Remarks"]
    assert "Hidden text found on page 1" in row["Remarks"]


def test_hidden_text_with_one_model_is_flagged_not_text_verified(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(tmp_path, monkeypatch, settings=ONE_MODEL, primary_a=[{"vendor": "Acme Trading", "total": "12.50"}])
    source = _write_pdf(tmp_path, MARKED_PAGES, hidden=["Invisible 99.99", ""])

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "SINGLE_MODEL"
    assert row["Models called"] == 1
    assert (
        "primary_a's answer could not be confirmed from the PDF text: the PDF has hidden text on page 1"
        in row["Remarks"]
    )


def test_hidden_text_stays_unverified_when_the_second_model_fails(tmp_path: Path, monkeypatch) -> None:
    processor, _ = _processor_with_providers(
        tmp_path,
        monkeypatch,
        settings=TWO_PRIMARIES,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
        primary_b=[],  # no scripted answer left, so the request fails
    )
    source = _write_pdf(tmp_path, MARKED_PAGES, hidden=["Invisible 99.99", ""])

    output = processor.handle_claimed_pdf(source).output_path

    assert output is not None
    row = _document_row(output)
    assert row["Validation status"] == "SINGLE_MODEL"
    assert row["Models called"] == 2
    assert "Second model called because primary_a's answer could not be confirmed: the PDF has hidden text on page 1" in row["Remarks"]
    assert "primary_b failed; only primary_a answered" in row["Remarks"]
    assert "confirmed from the PDF text: the PDF has hidden text on page 1" in row["Remarks"]


def test_failed_single_model_sends_the_pdf_to_errors(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(tmp_path, monkeypatch, settings=ONE_MODEL, primary_a=[])
    source = _write_pdf(tmp_path, MARKED_PAGES)

    outcome = processor.handle_claimed_pdf(source)

    # primary_a has no scripted answer left, so the request fails.
    assert outcome.output_path is None
    assert "too few usable model results" in outcome.error
    assert (outcome.archive_dir / "failure.txt").is_file()
    assert outcome.archive_dir.parent == tmp_path / "errors"


def test_same_pdf_twice_is_flagged_as_duplicate_without_new_calls(tmp_path: Path, monkeypatch) -> None:
    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        primary_a=[{"vendor": "Acme Trading", "total": "12.50"}],
    )
    source = _write_pdf(tmp_path, MARKED_PAGES)
    data = source.read_bytes()
    first = processor.handle_claimed_pdf(source)
    source.write_bytes(data)
    second = processor.handle_claimed_pdf(source)

    assert _document_row(first.output_path)["Duplicate status"] == "NONE"
    assert _document_row(second.output_path)["Duplicate status"] == "DUPLICATE"
    assert providers["primary_a"].calls == 1


class PageKeyedProvider:
    """Thread-safe offline provider that answers each document by its first page.

    The answers do not depend on call order, so documents can be read in
    parallel. A page without an answer makes the request fail.
    """

    def __init__(self, model_name: str, answers: dict[int, dict]) -> None:
        self.answers = answers
        self.calls = 0
        self.parts: list[bytes] = []
        self.in_flight = 0
        self.peak = 0
        self._lock = threading.Lock()
        self.config = ModelConfig(name=model_name, provider="gemini", model="test-model", api_key_env="X_API_KEY")

    def generate_json(self, prompt: str, parts: list[bytes]) -> ProviderResponse:
        with self._lock:
            self.calls += 1
            self.in_flight += 1
            self.peak = max(self.peak, self.in_flight)
            self.parts.extend(parts)
        try:
            time.sleep(0.05)
            first_page = int(re.search(r"PDF pages covered: \[(\d+)", prompt).group(1))
            if first_page not in self.answers:
                raise ProviderError(f"{self.config.name}: ServerError (HTTP 503)")
            return ProviderResponse(data={"fields": self.answers[first_page]}, usage=USAGE)
        finally:
            with self._lock:
                self.in_flight -= 1


# One-page documents whose "Page 1 of 1" marker settles every boundary, with
# enough text for the text check.
ONE_PAGE_DOCUMENTS = [
    f"Vendor {n} Trading\nPage 1 of 1\nDelivery address and payment terms\nTotal due {n}.00" for n in range(1, 7)
]


def _answers(*pages: int) -> dict[int, dict]:
    return {n: {"vendor": f"Vendor {n} Trading", "total": f"{n}.00"} for n in pages}


def _keyed_processor(
    root: Path,
    monkeypatch,
    answers: dict[int, dict],
    workers: int,
    model: dict | None = None,
    progress=None,
    stop=None,
) -> tuple[BatchProcessor, PageKeyedProvider]:
    settings = {
        "models": [{**ONE_MODEL["models"][0], **(model or {})}],
        "limits": {"extraction_workers": workers},
    }
    processor = BatchProcessor(load_config(_write_config(root, **settings)), progress=progress, stop_check=stop)
    provider = PageKeyedProvider("primary_a", answers)
    monkeypatch.setattr(BatchProcessor, "_provider", lambda self, model: provider)
    return processor, provider


def _document_rows(output: Path) -> list[dict]:
    header, *rows = list(load_workbook(output, data_only=True)["Documents"].iter_rows(values_only=True))
    return [dict(zip(header, row)) for row in rows]


def test_documents_are_read_in_parallel_and_kept_in_order(tmp_path: Path, monkeypatch) -> None:
    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 2, 3, 4), workers=3)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:4])

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    rows = _document_rows(outcome.output_path)
    assert [(row["Page start"], row["vendor"], row["total"]) for row in rows] == [
        (n, f"Vendor {n} Trading", float(n)) for n in range(1, 5)
    ]
    assert {row["Validation status"] for row in rows} == {"TEXT_VERIFIED"}
    assert provider.peak == 3


def test_pdf_parts_are_cut_before_the_workers_start(tmp_path: Path, monkeypatch) -> None:
    import doc_pipeline.processor as processor_module

    threads = []
    original = processor_module.extract_pdf_pages

    def recording_extract(*args):
        threads.append(threading.current_thread().name)
        return original(*args)

    monkeypatch.setattr(processor_module, "extract_pdf_pages", recording_extract)
    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 2, 3), workers=3, model={"input": "pdf"})
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:3])

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    # PyMuPDF is not thread-safe: every page cut happens on the calling thread.
    assert threads == [threading.current_thread().name] * 3
    assert len(provider.parts) == 3 and all(part.startswith(b"%PDF") for part in provider.parts)


def test_a_document_without_an_answer_does_not_fail_the_pdf(tmp_path: Path, monkeypatch) -> None:
    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 3, 4), workers=2)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:4])
    data = source.read_bytes()

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    rows = _document_rows(outcome.output_path)
    assert [row["Validation status"] for row in rows] == ["TEXT_VERIFIED", "FAILED", "TEXT_VERIFIED", "TEXT_VERIFIED"]
    assert rows[1]["vendor"] is None and "No usable model answer" in rows[1]["Remarks"]
    run_rows = dict(load_workbook(outcome.output_path, data_only=True)["Run info"].iter_rows(values_only=True))
    assert run_rows["Documents with no usable answer"] == "1"
    # The incomplete result is not cached: the same file again is read again.
    source.write_bytes(data)
    again = processor.handle_claimed_pdf(source)
    assert provider.calls == 8
    assert [row["Duplicate status"] for row in _document_rows(again.output_path)].count("DUPLICATE") == 0


def test_three_failures_in_a_row_stop_the_rest_of_the_pdf(tmp_path: Path, monkeypatch) -> None:
    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 5, 6), workers=1)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS)

    outcome = processor.handle_claimed_pdf(source)

    assert outcome.error is None
    rows = _document_rows(outcome.output_path)
    assert [row["Validation status"] for row in rows] == ["TEXT_VERIFIED", *["FAILED"] * 5]
    assert provider.calls == 4  # documents 5 and 6 were never sent
    assert all("Not read: stopped after 3 documents in a row" in row["Remarks"] for row in rows[4:])


def test_a_stop_between_documents_propagates_and_the_rest_are_not_sent(tmp_path: Path, monkeypatch) -> None:
    def stop() -> None:
        if provider.calls >= 1:  # the first document may be read, the second may not
            raise RunStopped("the run was cancelled")

    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 2, 3, 4), workers=1, stop=stop)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:4])

    with pytest.raises(RunStopped, match="the run was cancelled"):
        processor.handle_claimed_pdf(source)

    assert provider.calls == 1  # documents 2-4 were never sent
    # A stop is not a failed file: the PDF was left alone, not archived.
    assert source.exists() and not list((tmp_path / "errors").iterdir())


def test_a_stop_between_the_models_of_one_document_propagates(tmp_path: Path, monkeypatch) -> None:
    read_first = [None]

    def stop() -> None:
        primary_a = read_first[0]
        if primary_a is not None and primary_a.calls >= 1:
            raise RunStopped("the run was cancelled")

    processor, providers = _processor_with_providers(
        tmp_path,
        monkeypatch,
        stop=stop,
        primary_a=[{"vendor": "Acme", "total": "12.50"}],
        primary_b=[{"vendor": "Acme", "total": "12.50"}],
        arbiter=[{"vendor": "Acme", "total": "12.50"}],
    )
    read_first[0] = providers["primary_a"]
    source = _write_pdf(tmp_path)  # unmarked pages, so the boundary model is asked too

    with pytest.raises(RunStopped, match="the run was cancelled"):
        processor.handle_claimed_pdf(source)

    # The stop landed between primary_a's and primary_b's requests.
    assert providers["primary_a"].calls == 1
    assert providers["primary_b"].calls == 0 and providers["arbiter"].calls == 0
    # A stop is not a failed file: the PDF was left alone, not archived.
    assert source.exists() and not list((tmp_path / "errors").iterdir())


def test_a_stop_between_page_chunks_propagates(tmp_path: Path, monkeypatch) -> None:
    answered = threading.Event()

    def stop() -> None:
        if answered.is_set():  # the first chunk's reply has arrived
            raise RunStopped("the run was cancelled")

    class ChunkStoppingProvider(ScriptedProvider):
        """Flags the stop once its first extraction reply is in."""

        def generate_json(self, prompt: str, parts: list[bytes]) -> ProviderResponse:
            response = super().generate_json(prompt, parts)
            if "boundaries between" not in prompt:
                answered.set()
            return response

    config = load_config(_write_config(tmp_path, **{**ONE_MODEL, "limits": {"extraction_batch_pages": 1}}))
    processor = BatchProcessor(config, stop_check=stop)
    providers = {
        "boundary": ScriptedProvider("boundary", []),
        "primary_a": ChunkStoppingProvider("primary_a", [{"vendor": "Acme Trading", "total": "12.50"}]),
    }
    monkeypatch.setattr(BatchProcessor, "_provider", lambda self, model: providers[model.name])
    source = _write_pdf(tmp_path)  # two unmarked pages: one document, one chunk per page

    with pytest.raises(RunStopped, match="the run was cancelled"):
        processor.handle_claimed_pdf(source)

    # The boundary request and the first chunk ran; the second chunk did not.
    assert providers["boundary"].calls == 1 and providers["primary_a"].calls == 1
    # A stop is not a failed file: the PDF was left alone, not archived.
    assert source.exists() and not list((tmp_path / "errors").iterdir())


def test_a_stop_while_the_last_document_is_in_flight_keeps_the_file(tmp_path: Path, monkeypatch) -> None:
    answered_last = threading.Event()

    def stop() -> None:
        if answered_last.is_set():
            raise RunStopped("the run was cancelled")

    processor, provider = _keyed_processor(tmp_path, monkeypatch, _answers(1, 2), workers=1, stop=stop)
    original = provider.generate_json

    def answer_and_flag(prompt: str, parts: list[bytes]) -> ProviderResponse:
        response = original(prompt, parts)
        if provider.calls == 2:  # the last document of the file was answered
            answered_last.set()
        return response

    provider.generate_json = answer_and_flag
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:2])

    outcome = processor.handle_claimed_pdf(source)

    # The stop can only land after the last reply, so the paid-for results
    # are kept instead of the file being discarded.
    assert outcome.error is None
    assert [row["Validation status"] for row in _document_rows(outcome.output_path)] == ["TEXT_VERIFIED"] * 2
    assert provider.calls == 2


def test_a_stop_between_rules_only_documents_propagates(tmp_path: Path) -> None:
    checks = [0]

    def stop() -> None:
        checks[0] += 1
        if checks[0] > 1:
            raise RunStopped("the run was cancelled")

    config = load_config(_write_config(tmp_path, models=[], boundary_model=None))
    processor = BatchProcessor(config, stop_check=stop)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:3])

    with pytest.raises(RunStopped, match="the run was cancelled"):
        processor.handle_claimed_pdf(source)

    assert checks[0] == 2 and source.exists()


def test_a_stop_during_boundary_detection_propagates_not_a_failed_file(tmp_path: Path, monkeypatch) -> None:
    class StoppingBoundaryProvider(ScriptedProvider):
        """The run was stopped while it waited to retry the boundary request."""

        def generate_json(self, prompt: str, parts: list[bytes]) -> ProviderResponse:
            raise RunStopped("the run was cancelled")

    processor = BatchProcessor(load_config(_write_config(tmp_path)))
    monkeypatch.setattr(BatchProcessor, "_provider", lambda self, model: StoppingBoundaryProvider("boundary", []))
    source = _write_pdf(tmp_path)  # unmarked pages, so the boundary model is asked

    with pytest.raises(RunStopped, match="the run was cancelled"):
        processor.handle_claimed_pdf(source)

    # A stop is not a failed file: the PDF was left alone, not archived.
    assert source.exists() and not list((tmp_path / "errors").iterdir())


def test_progress_reports_each_step_and_document(tmp_path: Path, monkeypatch) -> None:
    events = []
    processor, _ = _keyed_processor(tmp_path, monkeypatch, _answers(1, 2, 3), workers=2, progress=events.append)
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:3])

    processor.handle_claimed_pdf(source)

    steps = [(event.stage, event.done, event.total) for event in events]
    assert steps == [("reading", 0, 0), ("splitting", 0, 0), *[("extracting", n, 3) for n in range(4)]]


def test_usage_is_counted_per_reply_and_cached_files_add_none(tmp_path: Path, monkeypatch) -> None:
    from doc_pipeline.providers import GeminiProvider

    def answer(self, prompt: str, parts) -> tuple[str, dict]:
        first_page = int(re.search(r"PDF pages covered: \[(\d+)", prompt).group(1))
        fields = {"vendor": f"Vendor {first_page} Trading", "total": f"{first_page}.00"}
        return json.dumps({"fields": fields}), USAGE

    monkeypatch.setattr(GeminiProvider, "_generate_text", answer)
    config = load_config(_write_config(tmp_path, **ONE_MODEL))
    # Real providers (not a patched _provider), so the reply hooks run.
    processor = BatchProcessor(config, api_keys={"primary_a": "key"})
    source = _write_pdf(tmp_path, ONE_PAGE_DOCUMENTS[:2])
    data = source.read_bytes()

    processor.handle_claimed_pdf(source)
    source.write_bytes(data)
    again = processor.handle_claimed_pdf(source)

    assert {row["Duplicate status"] for row in _document_rows(again.output_path)} == {"DUPLICATE"}
    # Two documents, one request each; the cached second file made no requests.
    assert processor.usage_by_model() == {
        ("gemini", "test-model"): {**{key: value * 2 for key, value in USAGE.items()}, "calls": 2}
    }
