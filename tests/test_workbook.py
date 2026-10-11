from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from doc_pipeline.models import BatchResult, DocumentResult, FieldSpec, ModelResult
from doc_pipeline.workbook import write_workbook


def test_workbook_writes_separate_validation_and_duplicate_flags(tmp_path: Path) -> None:
    usage_call = {"prompt_tokens": 100, "output_tokens": 20, "thoughts_tokens": 5, "cached_tokens": 0, "total_tokens": 120}
    batch = BatchResult(
        source_name="source.pdf",
        source_sha256="a" * 64,
        config_fingerprint="b" * 64,
        created_at="2026-01-01T00:00:00+00:00",
        documents=[
            DocumentResult(
                index=1,
                page_start=1,
                page_end=2,
                fields={"vendor": "=1+1", "total": "12.50"},
                validation_status="CONFLICT",
                duplicate_status="DUPLICATE",
                duplicate_of="a" * 64,
                model_results=[
                    ModelResult(
                        name="primary_a",
                        fields={"vendor": "=1+1", "total": "12.50"},
                        audit={"usage_calls": [usage_call], "usage": {**usage_call, "calls": 1}},
                        provider="gemini",
                        model_id="=test-model",
                    )
                ],
                remarks=["Starts at page 1: first page", "=cmd"],
            )
        ],
        boundary_audit=[
            {"model": "boundary", "provider": "gemini", "model_id": "test-model", "usage": usage_call}
        ],
        page_decisions=[
            {"page": 1, "decision": "start", "decided_by": "rule", "reason": "first page", "text_chars": 40, "blank": False},
            {"page": 2, "decision": "continuation", "decided_by": "model", "reason": "boundary model: same document", "confidence": 0.9},
        ],
    )
    path = write_workbook([batch], (FieldSpec("vendor"), FieldSpec("total", "decimal")), tmp_path / "result.xlsx")

    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Documents"]
    header = [cell.value for cell in sheet[1]]
    cells = {name: sheet.cell(row=2, column=index + 1) for index, name in enumerate(header)}
    assert cells["Validation status"].value == "CONFLICT"
    assert cells["Validation status"].fill.fgColor.rgb.endswith("F8D7DA")
    assert cells["Duplicate status"].value == "DUPLICATE"
    assert cells["Models called"].value == 1
    assert cells["Remarks"].value == "Starts at page 1: first page | =cmd"
    assert cells["vendor"].value == "'=1+1"
    assert cells["total"].value == 12.5

    pages = list(workbook["Pages"].iter_rows(values_only=True))
    assert pages[1][:4] == ("source.pdf", 1, "start", "rule")
    assert pages[2][:4] == ("source.pdf", 2, "continuation", "model")

    usage = workbook["Token usage"]
    rows = list(usage.iter_rows(values_only=True))
    header, boundary_row, extraction_row = rows[0], rows[1], rows[2]
    columns = {name: index for index, name in enumerate(header)}
    assert boundary_row[columns["Stage"]] == "boundary"
    assert extraction_row[columns["Stage"]] == "extraction"
    assert extraction_row[columns["Document #"]] == 1
    assert extraction_row[columns["Prompt tokens"]] == 100
    assert extraction_row[columns["Model ID"]] == "'=test-model"

    run_rows = dict((row[0], row[1]) for row in workbook["Run info"].iter_rows(values_only=True))
    assert "Tokens total" in run_rows
    assert "prompt=200" in run_rows["Tokens total"]
    assert run_rows["Pages decided by local rules"] == "1 of 2"
    assert run_rows["Pages sent to the boundary model"] == "1"
    # Filters on data tables only, not on the Run info summary.
    assert sheet.auto_filter.ref == sheet.dimensions
    assert workbook["Run info"].auto_filter.ref is None


def test_combined_workbook_lists_every_source_and_failure(tmp_path: Path) -> None:
    def batch(name: str, total: str) -> BatchResult:
        return BatchResult(
            source_name=name,
            source_sha256=name[0] * 64,
            config_fingerprint="c" * 64,
            created_at="2026-01-01T00:00:00+00:00",
            documents=[DocumentResult(index=1, page_start=1, page_end=1, fields={"=cmd": total}, validation_status="RULES_ONLY")],
            mode="Rules only (no AI model); page splitting: rules only",
        )

    path = write_workbook(
        [batch("a.pdf", "1.00"), batch("b.pdf", "2.00")],
        (FieldSpec("=cmd", "decimal"),),
        tmp_path / "combined.xlsx",
        failures=[("broken.pdf", "The file is not a readable PDF.")],
    )

    workbook = load_workbook(path)
    documents = list(workbook["Documents"].iter_rows(values_only=True))
    assert documents[0][-1] == "'=cmd"
    assert [(row[0], row[-1]) for row in documents[1:]] == [("a.pdf", 1.0), ("b.pdf", 2.0)]
    assert workbook["Documents"]["E2"].fill.fgColor.rgb.endswith("FFF3CD")
    assert list(workbook["Failed files"].iter_rows(values_only=True))[1] == ("broken.pdf", "The file is not a readable PDF.")
    run_rows = list(workbook["Run info"].iter_rows(values_only=True))
    assert ("Files processed", "2") in run_rows and ("Files failed", "1") in run_rows
    assert [value for label, value in run_rows if label == "Source PDF"] == ["a.pdf", "b.pdf"]
    # No model calls: Token usage holds only its header row, so it gets no filter.
    assert workbook["Token usage"].max_row == 1 and workbook["Token usage"].auto_filter.ref is None


def test_documents_without_an_answer_are_marked_as_problems(tmp_path: Path) -> None:
    batch = BatchResult(
        source_name="a.pdf",
        source_sha256="a" * 64,
        config_fingerprint="c" * 64,
        created_at="2026-01-01T00:00:00+00:00",
        documents=[DocumentResult(index=1, page_start=1, page_end=1, fields={"total": None}, validation_status="FAILED")],
    )

    workbook = load_workbook(write_workbook([batch], (FieldSpec("total", "decimal"),), tmp_path / "failed.xlsx"))

    assert workbook["Documents"]["E2"].value == "FAILED"
    assert workbook["Documents"]["E2"].fill.fgColor.rgb.endswith("F8D7DA")
    run_rows = dict(workbook["Run info"].iter_rows(values_only=True))
    assert run_rows["Documents with no usable answer"] == "1"
