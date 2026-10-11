from __future__ import annotations

import os
import tempfile
from collections import Counter
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import BatchResult, FieldSpec
from .schema import audit_json, excel_safe_string, excel_value

USAGE_KEYS = ("prompt_tokens", "output_tokens", "thoughts_tokens", "cached_tokens", "total_tokens")
PROBLEM_FILL = PatternFill("solid", fgColor="F8D7DA")
REVIEW_FILL = PatternFill("solid", fgColor="FFF3CD")
STATUS_FILLS = {
    "CONFLICT": PROBLEM_FILL,
    "FAILED": PROBLEM_FILL,
    "INCOMPLETE_VALIDATION": REVIEW_FILL,
    "CONSENSUS_2_OF_3": REVIEW_FILL,
    "SINGLE_MODEL": REVIEW_FILL,
    "RULES_ONLY": REVIEW_FILL,
    "DUPLICATE": REVIEW_FILL,
    "POSSIBLE_DUPLICATE": REVIEW_FILL,
}
UsageRow = tuple[str, str, str | int, str, str, str, dict[str, int]]


def _usage_rows(batches: Sequence[BatchResult]) -> list[UsageRow]:
    rows: list[UsageRow] = []
    for batch in batches:
        for boundary in batch.boundary_audit:
            usage = boundary.get("usage") or {}
            if usage:
                rows.append(
                    (
                        batch.source_name,
                        "boundary",
                        "-",
                        str(boundary.get("model", "boundary")),
                        str(boundary.get("provider", "")),
                        str(boundary.get("model_id", "")),
                        {key: int(usage.get(key, 0) or 0) for key in USAGE_KEYS},
                    )
                )
        for document in batch.documents:
            for result in document.model_results:
                for call in result.audit.get("usage_calls") or []:
                    rows.append(
                        (
                            batch.source_name,
                            "extraction",
                            document.index,
                            result.name,
                            result.provider or "",
                            result.model_id or "",
                            {key: int(call.get(key, 0) or 0) for key in USAGE_KEYS},
                        )
                    )
    return rows


def _usage_totals(rows: list[UsageRow]) -> dict[str, dict[str, int]]:
    totals: dict[str, dict[str, int]] = {}
    for _, _, _, name, provider, model_id, usage in rows:
        label = f"{name} ({provider}/{model_id})" if provider or model_id else name
        bucket = totals.setdefault(label, {key: 0 for key in USAGE_KEYS} | {"calls": 0})
        for key in USAGE_KEYS:
            bucket[key] += usage[key]
        bucket["calls"] += 1
    return totals


def _usage_summary(bucket: dict[str, int]) -> str:
    return (
        f"prompt={bucket['prompt_tokens']} output={bucket['output_tokens']} "
        f"thoughts={bucket['thoughts_tokens']} cached={bucket['cached_tokens']} "
        f"total={bucket['total_tokens']} calls={bucket['calls']}"
    )


def _call_summary(batch: BatchResult) -> list[tuple[str, str | int]]:
    """How many pages and documents needed a model, for the Run info sheet."""
    decided_by = Counter(str(page.get("decided_by") or "") for page in batch.page_decisions)
    return [
        ("Pages decided by local rules", f"{decided_by['rule']} of {len(batch.page_decisions)}"),
        ("Pages sent to the boundary model", decided_by["model"]),
        ("Pages guessed without a model", decided_by["guess"]),
        ("Boundary model requests", len(batch.boundary_audit)),
        ("Extraction model runs", sum(len(document.model_results) for document in batch.documents)),
        (
            "Documents confirmed by the PDF text (one model)",
            sum(document.validation_status == "TEXT_VERIFIED" for document in batch.documents),
        ),
        ("Documents with no usable answer", sum(document.validation_status == "FAILED" for document in batch.documents)),
    ]


def write_workbook(
    batches: Sequence[BatchResult],
    field_specs: tuple[FieldSpec, ...],
    output_path: Path,
    failures: Sequence[tuple[str, str]] = (),
) -> Path:
    """Write one workbook for one or more source PDFs.

    failures lists (source PDF name, reason) for files that went to errors.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    documents_sheet = workbook.active
    documents_sheet.title = "Documents"
    headers = [
        "Source PDF",
        "Document #",
        "Page start",
        "Page end",
        "Validation status",
        "Duplicate status",
        "Duplicate of",
        "Field issues",
        "Models called",
        "Remarks",
        *[excel_safe_string(spec.name) for spec in field_specs],
    ]
    documents_sheet.append(headers)
    for batch in batches:
        for document in batch.documents:
            row: list[Any] = [
                excel_safe_string(batch.source_name),
                document.index,
                document.page_start,
                document.page_end,
                document.validation_status,
                document.duplicate_status,
                document.duplicate_of,
                audit_json(document.field_issues),
                len(document.model_results),
                excel_safe_string(" | ".join(document.remarks)),
            ]
            row.extend(excel_value(document.fields.get(spec.name), spec) for spec in field_specs)
            documents_sheet.append(row)
            cells = documents_sheet[documents_sheet.max_row]
            for cell in (cells[4], cells[5]):
                fill = STATUS_FILLS.get(str(cell.value))
                if fill:
                    cell.fill = fill
            if document.field_issues:
                cells[7].fill = REVIEW_FILL

    audit_sheet = workbook.create_sheet("Model audit")
    audit_sheet.append(["Source PDF", "Document #", "Pages", "Model", "Result", "Fields / structured response"])
    for batch in batches:
        source = excel_safe_string(batch.source_name)
        for boundary in batch.boundary_audit:
            audit_sheet.append(
                [
                    source,
                    "-",
                    "boundary pages",
                    excel_safe_string(
                        f"{boundary.get('model', 'boundary')} ({boundary.get('provider', '')}/{boundary.get('model_id', '')})"
                    ),
                    "BOUNDARY_DETECTION",
                    audit_json(boundary),
                ]
            )
        for document in batch.documents:
            pages = f"{document.page_start}-{document.page_end}"
            if not document.model_results:
                audit_sheet.append(
                    [source, document.index, pages, "rules (no AI model)", document.validation_status, audit_json(document.remarks)]
                )
            for result in document.model_results:
                audit_sheet.append(
                    [
                        source,
                        document.index,
                        pages,
                        excel_safe_string(f"{result.name} ({result.provider or ''}/{result.model_id or ''})"),
                        excel_safe_string(result.error or "OK"),
                        audit_json({"fields": result.fields, "audit": result.audit}),
                    ]
                )

    pages_sheet = workbook.create_sheet("Pages")
    pages_sheet.append(
        ["Source PDF", "Page", "Decision", "Decided by", "Reason", "Confidence", "Page marker", "Titles", "Text characters", "Blank"]
    )
    for batch in batches:
        for page in batch.page_decisions:
            pages_sheet.append(
                [
                    excel_safe_string(batch.source_name),
                    page.get("page"),
                    page.get("decision"),
                    page.get("decided_by") or "",
                    excel_safe_string(str(page.get("reason") or "")),
                    page.get("confidence"),
                    excel_safe_string(str(page.get("page_marker") or "")),
                    excel_safe_string(str(page.get("titles") or "")),
                    page.get("text_chars"),
                    "yes" if page.get("blank") else "no",
                ]
            )

    usage_sheet = workbook.create_sheet("Token usage")
    usage_sheet.append(
        ["Source PDF", "Stage", "Document #", "Model", "Provider", "Model ID", *[key.replace("_", " ").capitalize() for key in USAGE_KEYS]]
    )
    usage_rows = _usage_rows(batches)
    for source_name, stage, document_index, name, provider, model_id, usage in usage_rows:
        usage_sheet.append(
            [
                excel_safe_string(source_name),
                stage,
                document_index,
                *[excel_safe_string(text) for text in (name, provider, model_id)],
                *[usage[key] for key in USAGE_KEYS],
            ]
        )

    if failures:
        failed_sheet = workbook.create_sheet("Failed files")
        failed_sheet.append(["Source PDF", "Reason"])
        for source_name, reason in failures:
            failed_sheet.append([excel_safe_string(source_name), excel_safe_string(reason)])

    run_sheet = workbook.create_sheet("Run info")
    run_sheet.append(["Item", "Value"])
    run_rows: list[tuple[str, Any]] = []
    if len(batches) > 1 or failures:
        run_rows += [("Files processed", len(batches)), ("Files failed", len(failures))]
    for batch in batches:
        run_rows += [
            ("Source PDF", batch.source_name),
            ("Mode", batch.mode or "-"),
            ("Source SHA-256", batch.source_sha256),
            ("Configuration fingerprint", batch.config_fingerprint),
            ("Created (UTC)", batch.created_at),
            ("Documents found", len(batch.documents)),
            *_call_summary(batch),
        ]
    for label, value in run_rows:
        run_sheet.append([excel_safe_string(label), excel_safe_string(str(value))])
    totals = _usage_totals(usage_rows)
    for label, bucket in totals.items():
        run_sheet.append([excel_safe_string(f"Tokens {label}"), excel_safe_string(_usage_summary(bucket))])
    if usage_rows:
        grand = {key: sum(bucket[key] for bucket in totals.values()) for key in USAGE_KEYS}
        grand["calls"] = len(usage_rows)
        run_sheet.append([excel_safe_string("Tokens total"), excel_safe_string(_usage_summary(grand))])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        # Filters only on data tables with rows, not on the Run info summary.
        if sheet.max_row > 1 and sheet.title != "Run info":
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25476A")
        for column_cells in sheet.columns:
            column_index = column_cells[0].column
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 48)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

    handle, temporary_name = tempfile.mkstemp(prefix=f".{output_path.stem}-", suffix=".tmp.xlsx", dir=output_path.parent)
    os.close(handle)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        with temporary_path.open("rb+") as stream:
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary_path, output_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
    return output_path
